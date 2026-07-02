import json
import math
import re
import subprocess
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
INDEX = ROOT / "index.html"
OUT_DIR = ROOT / "proposal_output"
NODE = Path("/Users/minmacbook/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node")
PACKAGE_DISCOUNT_RATE = 0.15

ACCENT = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
PRINT_BLUE = RGBColor(0, 92, 185)
DARK = RGBColor(34, 34, 34)
MUTED = RGBColor(102, 102, 102)
GREEN = RGBColor(38, 112, 74)
RED = RGBColor(156, 44, 44)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF4EF"
PALE_RED = "F8EDEE"
PALE_PRINT_BLUE = "D7E9FF"

CARE_OPTIONS = {
    "serviceFree": {
        "label": "서비스프리",
        "short": "서비스프리",
        "alt": "절약형",
        "summary": "방문관리 비용 부담을 낮추고 딥클리닝 쿠폰 혜택을 제공하는 절약형 옵션",
        "headline": "7월 프로모션 15% 할인 기준으로 일시불 합산가와 5년 총렌탈료를 비교합니다.",
        "talk": "고객님, 서비스프리는 방문관리 비용 부담을 줄이고 필요한 혜택만 남긴 절약형 옵션입니다. 7월 프로모션 15% 할인 기준으로 매트리스와 파운데이션을 함께 렌탈하실 때의 5년 총렌탈료를 일시불 합산가와 비교해 보실 수 있습니다.",
    },
    "basic": {
        "label": "베이직케어",
        "short": "베이직",
        "alt": "깔끔형",
        "summary": "4개월 주기 전문가 방문 관리가 포함된 위생관리형 옵션",
        "headline": "렌탈 총액은 올라가지만, 방문관리까지 포함한 관리형 선택지입니다.",
        "talk": "고객님, 베이직케어는 단순히 제품 가격만 비교하면 일시불보다 총렌탈료가 높아질 수 있습니다. 대신 5년 동안 4개월 주기의 전문가 방문 관리가 포함되기 때문에, 제품 구매가 아니라 관리까지 포함한 이용 방식으로 보시면 됩니다.",
    },
    "special": {
        "label": "스페셜체인지",
        "short": "스페셜",
        "alt": "실속형",
        "summary": "새 탑퍼 또는 커버 교체 혜택과 딥클리닝 쿠폰이 포함된 교체형 옵션",
        "headline": "탑퍼/커버 교체 혜택을 비용으로 포함해 보는 실속형 비교표입니다.",
        "talk": "고객님, 스페셜체인지는 서비스프리보다 월 렌탈료가 올라가지만, 사용 중 새 탑퍼 또는 커버 교체 혜택이 포함됩니다. 일시불과 단순 가격만 비교하기보다는, 교체 혜택까지 포함한 5년 이용가로 비교해 보시면 좋습니다.",
    },
    "total": {
        "label": "토탈케어",
        "short": "토탈",
        "alt": "완벽형",
        "summary": "4개월 방문관리와 새 탑퍼 또는 커버 교체를 함께 포함한 완전관리형 옵션",
        "headline": "가장 높은 관리 혜택을 포함한 프리미엄 이용가 비교표입니다.",
        "talk": "고객님, 토탈케어는 방문관리와 탑퍼 또는 커버 교체 혜택을 모두 포함하기 때문에 총렌탈료가 가장 높게 나올 수 있습니다. 대신 5년 동안 제품 사용과 관리 부담을 함께 줄이는 프리미엄 관리형 선택지입니다.",
    },
}

MODEL_ADD_FEES = {
    "hybrid4": 5000,
    "lunaire": 4000,
    "modi": 4000,
    "elite": 1000,
    "smarts8": 4000,
    "sigsc": 6000,
    "doubleside": 2000,
    "doublechainge": 5000,
    "compactfoam": 2000,
}
NO_TOPPER_MODELS = {"sigs"}
DEFAULT_CARE_FEE = 2000
DEFAULT_MBC = {
    "serviceFree": {3: 4000, 5: 4000, 7: 4000, 9: 4000},
    "special": {3: 4000, 5: 4000, 7: 4000, 9: 4000},
    "basic": {3: 0, 5: 0, 7: 0, 9: 0},
    "total": {3: 0, 5: 0, 7: 0, 9: 0},
}


def extract_js_object(source, const_name):
    marker = f"const {const_name} ="
    start = source.index(marker)
    brace_start = source.index("{", start)
    depth = 0
    in_str = None
    escape = False
    for i in range(brace_start, len(source)):
        ch = source[i]
        if in_str:
            if escape:
                escape = False
            elif ch == "\\":
                escape = True
            elif ch == in_str:
                in_str = None
            continue
        if ch in ("'", '"', "`"):
            in_str = ch
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return source[brace_start : i + 1]
    raise RuntimeError(f"Could not extract {const_name}")


def load_product_data():
    source = INDEX.read_text(encoding="utf-8")
    obj = extract_js_object(source, "PRODUCT_DATA")
    code = f"const PRODUCT_DATA = {obj}; process.stdout.write(JSON.stringify(PRODUCT_DATA));"
    result = subprocess.run([str(NODE), "-e", code], check=True, capture_output=True, text=True)
    return json.loads(result.stdout)


def ceil10(value):
    return int(math.ceil(value / 10.0) * 10)


def money(value):
    return f"{int(round(value)):,}원"


def manwon(value):
    return f"{abs(value) / 10000:,.1f}만원"


def diff_text(value):
    if value > 0:
        return f"{manwon(value)} 낮음"
    if value < 0:
        return f"{manwon(value)} 높음"
    return "동일"


def effective_care(care_id, m_key):
    if m_key not in NO_TOPPER_MODELS:
        return care_id
    if care_id == "total":
        return "basic"
    if care_id == "special":
        return "serviceFree"
    return care_id


def pricing(care_id, m_key, m_item, f_item):
    period = 5
    months = 60
    discount_rate = PACKAGE_DISCOUNT_RATE

    raw_monthly = m_item["rentalData"][str(period)]["monthly"]
    eff_care = effective_care(care_id, m_key)
    c_fee = DEFAULT_CARE_FEE if eff_care in ("basic", "total") else 0
    a_fee = MODEL_ADD_FEES.get(m_key, 0) if eff_care in ("special", "total") else 0
    m_bc = DEFAULT_MBC.get(care_id, {}).get(period, 0)
    m_le = 5000 if m_key == "lunaire" else 0
    m_se = 9000 if m_key == "smarts8" else 0
    sf_mbc = DEFAULT_MBC["serviceFree"][period]

    m_base = raw_monthly - sf_mbc + c_fee + a_fee + m_le + m_se
    m_pre_pkg = m_base - 1000 - m_bc - m_le - m_se
    m_pkg_d = ceil10(m_pre_pkg * discount_rate)
    m_final = max(0, m_pre_pkg - m_pkg_d)

    f_raw = f_item["rentalData"][str(period)]["monthly"]
    f_pkg_d = ceil10(f_raw * discount_rate)
    f_final = max(0, f_raw - f_pkg_d)

    monthly = m_final + f_final
    rent_total = monthly * months
    lump_total = int(m_item["price"]) + int(f_item["price"])
    diff = lump_total - rent_total
    return {
        "effective_care": eff_care,
        "m_monthly": m_final,
        "f_monthly": f_final,
        "monthly": monthly,
        "rent_total": rent_total,
        "lump_total": lump_total,
        "diff": diff,
        "diff_rate": diff / lump_total if lump_total else 0,
        "m_model_no": m_item.get("modelNo", ""),
        "f_model_no": f_item.get("modelNo", ""),
    }


def build_rows(data, care_id):
    foundation_items = {
        item["size"]: item
        for item in data["frames"]["foundation"]["items"]
        if "5" in item.get("rentalData", {})
    }
    rows = []
    for m_key, mattress in data["mattresses"].items():
        for item in mattress["items"]:
            size = item["size"]
            if size not in foundation_items or "5" not in item.get("rentalData", {}):
                continue
            rows.append(
                {
                    "m_key": m_key,
                    "mattress": mattress["name"],
                    "size": size,
                    "m_price": int(item["price"]),
                    **pricing(care_id, m_key, item, foundation_items[size]),
                }
            )
    size_rank = {"슈싱": 0, "퀸": 1, "킹": 2, "라지킹": 3, "그레이트킹": 4}
    return sorted(rows, key=lambda r: (size_rank.get(r["size"], 99), r["m_price"], r["mattress"]))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C6D6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, color=None, bold=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def apply_section_layout(section, landscape=False, compact=False):
    if landscape:
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        margin = 0.28 if compact else 0.5
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        margin = 0.75
    section.top_margin = Inches(margin)
    section.bottom_margin = Inches(margin)
    section.left_margin = Inches(0.35 if compact else 0.78)
    section.right_margin = Inches(0.35 if compact else 0.78)
    section.header_distance = Inches(0.16 if compact else 0.35)
    section.footer_distance = Inches(0.14 if compact else 0.35)


def add_para(doc, text="", size=11, color=DARK, bold=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12.5, color=BLUE if level == 1 else ACCENT, bold=True)
    return p


def add_metric_table(doc, metrics):
    table = doc.add_table(rows=1, cols=len(metrics))
    set_table_borders(table, color="D5DEE8", size="4")
    set_table_widths(table, [6.5 / len(metrics)] * len(metrics))
    for idx, (label, value, note, good) in enumerate(metrics):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, PALE_GREEN if good else LIGHT_GRAY)
        set_cell_margins(cell, top=130, bottom=130, start=140, end=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=9, color=MUTED, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(value)
        set_run_font(r2, size=15, color=GREEN if good else ACCENT, bold=True)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(note)
        set_run_font(r3, size=8.4, color=MUTED)
    return table


def add_simple_table(doc, headers, rows, widths, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    set_table_widths(table, widths)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell, top=85, bottom=85, start=90, end=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=8.2, color=ACCENT, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell, top=70, bottom=70, start=85, end=85)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx in (0, 1) else WD_ALIGN_PARAGRAPH.RIGHT
            text = str(val)
            r = p.add_run(text)
            is_diff = idx == len(row) - 1
            color = RED if is_diff and "높음" in text else GREEN if is_diff and "낮음" in text else DARK
            set_run_font(r, size=font_size, color=color, bold=is_diff)
    return table


def add_readable_compare_table(doc, headers, rows, widths, highlight_lower=False, font_size=10.2, header_size=10.4, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table, color="A8BCD4", size="5")
    set_table_widths(table, widths)
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        if compact:
            set_cell_margins(cell, top=64, bottom=64, start=76, end=76)
        else:
            set_cell_margins(cell, top=78, bottom=78, start=85, end=85)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(header)
        set_run_font(r, size=header_size, color=ACCENT, bold=True)

    for row in rows:
        lower = bool(row.get("_lower"))
        values = row["values"]
        cells = table.add_row().cells
        for idx, val in enumerate(values):
            cell = cells[idx]
            if highlight_lower and lower:
                set_cell_shading(cell, PALE_PRINT_BLUE)
            if compact:
                set_cell_margins(cell, top=42, bottom=42, start=72, end=72)
            else:
                set_cell_margins(cell, top=58, bottom=58, start=80, end=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            if idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(val))
            is_diff = idx == len(values) - 1
            color = PRINT_BLUE if highlight_lower and lower else DARK
            if is_diff and not (highlight_lower and lower):
                color = RED if "높음" in str(val) else GREEN if "낮음" in str(val) else DARK
            set_run_font(r, size=font_size, color=color, bold=(is_diff or (highlight_lower and lower)))
    return table


def add_detail_table(doc, rows, care_id, font_size=10.8):
    headers = ["매트리스", "사이즈", "매트 월", "파데 월", "월 합계", "5년 총렌탈료", "일시불 총액", "차액"]
    widths = [1.43, 0.56, 0.82, 0.82, 0.86, 1.16, 1.16, 0.88]
    table_rows = []
    for r in rows:
        table_rows.append(
            {
                "_lower": r["diff"] > 0,
                "values": [
                    r["mattress"],
                    r["size"],
                    money(r["m_monthly"]),
                    money(r["f_monthly"]),
                    money(r["monthly"]),
                    money(r["rent_total"]),
                    money(r["lump_total"]),
                    diff_text(r["diff"]),
                ],
            }
        )
    return add_readable_compare_table(
        doc,
        headers,
        table_rows,
        widths,
        highlight_lower=True,
        font_size=font_size,
        header_size=10.8,
        compact=True,
    )


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("• ")
    set_run_font(r, size=10.2, color=ACCENT, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.2, color=DARK)


def add_callout(doc, title, body, positive=True):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="C7D8C7" if positive else "E2B8BA", size="6")
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GREEN if positive else PALE_RED)
    set_cell_margins(cell, top=125, bottom=125, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, size=10.8, color=GREEN if positive else RED, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.1, color=DARK)


def set_document_styles(doc, care):
    section = doc.sections[0]
    apply_section_layout(section)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    header = section.header.paragraphs[0]
    header.text = f"비렉스 5년 {care['label']} 세트 비교 제안"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.3, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = "앱 등록 가격 기준 · 고객 안내 및 현장 상담용"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, size=8.3, color=MUTED)


def summarize(rows):
    higher = [r for r in rows if r["diff"] < 0]
    lower = [r for r in rows if r["diff"] > 0]
    avg_diff = sum(r["diff"] for r in rows) / len(rows)
    return {
        "higher_count": len(higher),
        "lower_count": len(lower),
        "same_count": len(rows) - len(higher) - len(lower),
        "avg_diff": avg_diff,
        "max_higher": min(rows, key=lambda r: r["diff"]),
        "max_lower": max(rows, key=lambda r: r["diff"]),
        "exceptions": sum(1 for r in rows if r["effective_care"] != rows[0].get("requested_care", r["effective_care"])),
    }


def add_title_page(doc, care_id, rows, summary):
    care = CARE_OPTIONS[care_id]
    add_para(doc, "CUSTOMER PROPOSAL", size=9, color=MUTED, bold=True, after=3, before=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(f"5년 {care['label']}\n매트리스+파운데이션 비용 비교 제안")
    set_run_font(run, size=22, color=ACCENT, bold=True)
    add_para(
        doc,
        f"일시불 합산가와 5년 총렌탈료를 같은 기준으로 놓고 보는 {care['alt']} 상담 자료",
        size=11.5,
        color=MUTED,
        after=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_metric_table(
        doc,
        [
            ("비교 조합", f"{len(rows)}개", "매트리스 × 파운데이션", False),
            ("렌탈 총액 높음", f"{summary['higher_count']}개", "일시불 합산가 대비", False),
            ("렌탈 총액 낮음", f"{summary['lower_count']}개", "일시불보다 낮은 조합", summary["lower_count"] > 0),
        ],
    )
    add_para(doc, "", after=4)
    add_callout(doc, f"{care['label']} 제안의 핵심", care["headline"], positive=summary["lower_count"] > 0)
    add_heading(doc, "비교 기준", level=1)
    for text in [
        f"약정과 옵션: 5년 약정, {care['label']}, 매트리스 1대 + 파운데이션 1대.",
        "총렌탈료: 7월 프로모션 15% 할인 적용 후 월 렌탈료 60개월 합계.",
        "월 렌탈료: 신규 2개 세트 기준 앱 패키지 할인 15%와 케어 옵션별 추가 비용을 반영.",
        "일시불 총액: 앱 등록 매트리스 일시불가와 파운데이션 일시불가의 단순 합산.",
    ]:
        add_bullet(doc, text)
    add_para(doc, f"산출일: {date.today().isoformat()} · 출처: 현재 앱 index.html 등록 가격 및 계산식", size=9, color=MUTED, after=0)


def add_insights(doc, care_id, rows, summary):
    care = CARE_OPTIONS[care_id]
    add_heading(doc, "현장 설명 포인트", level=1)
    max_higher = summary["max_higher"]
    max_lower = summary["max_lower"]
    texts = [
        f"{care['label']}은 {care['summary']}입니다.",
        f"전체 {len(rows)}개 조합 중 {summary['higher_count']}개는 5년 총렌탈료가 일시불 합산가보다 높고, {summary['lower_count']}개는 낮게 산출됩니다.",
    ]
    if care_id == "serviceFree":
        texts.append("서비스프리는 전체 조합에서 렌탈 총액이 낮게 산출되므로, 상세표의 모든 행이 파란색 배경으로 표시됩니다.")
    else:
        texts.extend(
            [
                f"렌탈 총액이 가장 높아지는 조합은 {max_higher['mattress']}+파운데이션 {max_higher['size']}이며, 일시불 대비 {manwon(max_higher['diff'])} 높습니다.",
                f"프로모션 효과가 가장 큰 조합은 {max_lower['mattress']}+파운데이션 {max_lower['size']}이며, 일시불 대비 {diff_text(max_lower['diff'])}입니다.",
                "파란색 배경은 일시불보다 5년 총렌탈료가 낮은 조합만 표시합니다.",
            ]
        )
    texts.append("따라서 이 표는 월 납입액만 보는 자료가 아니라, 고객이 5년 총액 기준으로 판단할 수 있게 돕는 자료입니다.")
    for text in texts:
        add_bullet(doc, text)
    if care_id in ("special", "total"):
        add_bullet(doc, "탑퍼 미포함 모델은 앱 계산식상 교체 혜택이 제외되어 낮은 케어 기준으로 산출됩니다.")


def add_representative_page(doc, care_id, rows):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section_layout(section)
    add_heading(doc, "대표 조합 비교", level=1)
    add_para(
        doc,
        "할인율이 높은 순서로 고객에게 먼저 보여주기 좋은 대표 조합입니다. 전체 판단은 3~4장 상세표를 기준으로 확인해 주세요.",
        size=11.2,
        color=MUTED,
        after=10,
    )
    representative = sorted(rows, key=lambda r: r["diff_rate"], reverse=True)[:10]
    table_rows = []
    for r in representative:
        table_rows.append(
            {
                "_lower": r["diff"] > 0,
                "values": [
                    f"{r['mattress']} + 파운데이션",
                    r["size"],
                    money(r["monthly"]),
                    money(r["rent_total"]),
                    money(r["lump_total"]),
                    f"{r['diff_rate'] * 100:.1f}%",
                    diff_text(r["diff"]),
                ],
            }
        )
    add_readable_compare_table(
        doc,
        ["조합", "사이즈", "월 렌탈료", "5년 총렌탈료", "일시불 총액", "할인율", "차액"],
        table_rows,
        [1.70, 0.55, 0.90, 1.12, 1.12, 0.62, 0.92],
        highlight_lower=True,
        font_size=10.0,
        header_size=10.2,
    )
    add_para(doc, "파란색 배경 행은 5년 총렌탈료가 일시불 합산가보다 낮은 조합입니다.", size=10.4, color=PRINT_BLUE, bold=True, before=8, after=0)


def add_detail_page(doc, care_id, rows, sizes, page_label, start_new_page=True):
    if start_new_page:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
    else:
        section = doc.sections[-1]
    apply_section_layout(section, landscape=True, compact=True)
    care = CARE_OPTIONS[care_id]
    filtered = [r for r in rows if r["size"] in sizes]
    add_para(
        doc,
        f"5년 {care['label']} 기준 전체 조합 상세표({page_label})",
        size=16.8,
        color=BLUE,
        bold=True,
        after=2,
    )
    add_para(
        doc,
        "매트리스+파운데이션 동시 렌탈 · 7월 프로모션 15% 할인 반영 · 5년 약정 기준",
        size=11.8,
        color=DARK,
        bold=True,
        after=2,
    )
    add_para(
        doc,
        "정렬: 슈싱 > 퀸 > 킹 > 라지킹 · 파란색 배경 = 5년 총렌탈료가 일시불보다 낮은 조합 · 차액 = 일시불 총액 - 5년 총렌탈료",
        size=9.8,
        color=MUTED,
        after=4,
    )
    add_detail_table(doc, filtered, care_id, font_size=10.8)


def add_customer_script(doc, care_id):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section_layout(section)
    care = CARE_OPTIONS[care_id]
    add_heading(doc, "현장 활용 멘트", level=1)
    add_callout(doc, "상담 시 권장 문장", f"“{care['talk']} 이 자료는 월 납입액만 보여드리는 표가 아니라, 60개월 총렌탈료와 일시불 합산가를 같은 기준으로 비교한 자료라 실제 부담 차이를 투명하게 확인하실 수 있습니다.”")
    add_heading(doc, "팀 공유용 설명", level=2)
    for text in [
        "이 자료는 고객에게 렌탈과 일시불의 총액 차이를 숨기지 않고 보여주기 위해 만들었습니다.",
        "서비스프리처럼 프로모션으로 총액이 낮아지는 옵션과, 관리 혜택 때문에 총액이 올라가는 케어 옵션을 구분해서 안내하는 데 목적이 있습니다.",
        "현장에서는 먼저 고객이 원하는 관리 수준을 확인하고, 그 다음 월 납입액이 아니라 5년 총액 기준으로 비교해 주세요.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "유의사항", level=2)
    for text in [
        "본 자료는 현재 앱에 등록된 가격과 7월 프로모션 15% 할인 계산식 기준이며, 본사 가격표·프로모션 변경 시 결과가 달라질 수 있습니다.",
        "고객의 기존 렌탈 보유 여부, 카드 청구할인, 재렌탈 조건, 추가 제품 결합 여부는 별도 비교가 필요합니다.",
        "스페셜체인지와 토탈케어는 모델별 탑퍼/커버 구조에 따라 제공 혜택이 달라질 수 있으므로 계약 전 최종 조건을 확인해야 합니다.",
    ]:
        add_bullet(doc, text)


def audit_docx(path, care):
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    required = [care["label"], "전체 조합 상세표", "현장 활용 멘트", "일시불 총액 - 5년 총렌탈료"]
    missing = [s for s in required if s not in xml]
    if missing:
        raise AssertionError(f"Missing expected text in {path.name}: {missing}")
    if re.search(r"TODO|TBD|PLACEHOLDER", xml):
        raise AssertionError(f"Placeholder text remains in {path.name}")


def build_doc(care_id, data):
    rows = build_rows(data, care_id)
    for row in rows:
        row["requested_care"] = care_id
    summary = summarize(rows)
    care = CARE_OPTIONS[care_id]
    doc = Document()
    set_document_styles(doc, care)
    add_title_page(doc, care_id, rows, summary)
    add_insights(doc, care_id, rows, summary)
    add_representative_page(doc, care_id, rows)
    add_detail_page(doc, care_id, rows, {"슈싱", "퀸"}, "슈싱, 퀸")
    add_detail_page(doc, care_id, rows, {"킹", "라지킹"}, "킹, 라지킹")
    add_customer_script(doc, care_id)
    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / f"비렉스_5년_{care['label']}_매트리스_파운데이션_비교제안서.docx"
    doc.save(path)
    audit_docx(path, care)
    return path, rows, summary


def set_detail_collection_styles(doc):
    section = doc.sections[0]
    apply_section_layout(section, landscape=True, compact=True)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.8)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    header = section.header.paragraphs[0]
    header.text = "비렉스 5년 케어 옵션 전체조합 상세표 모음"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.3, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = "앱 등록 가격 기준 · 상세표 단독 인쇄용"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, size=8.3, color=MUTED)


def audit_detail_collection(path):
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    required = [
        "5년 서비스프리 기준 전체 조합 상세표",
        "5년 베이직케어 기준 전체 조합 상세표",
        "5년 스페셜체인지 기준 전체 조합 상세표",
        "5년 토탈케어 기준 전체 조합 상세표",
        "7월 프로모션 15% 할인 반영",
        "파란색 배경 = 5년 총렌탈료가 일시불보다 낮은 조합",
    ]
    missing = [s for s in required if s not in xml]
    if missing:
        raise AssertionError(f"Missing expected text in {path.name}: {missing}")
    if re.search(r"TODO|TBD|PLACEHOLDER", xml):
        raise AssertionError(f"Placeholder text remains in {path.name}")


def build_detail_collection_doc(data):
    doc = Document()
    set_detail_collection_styles(doc)
    first_page = True
    for care_id in ("serviceFree", "basic", "special", "total"):
        rows = build_rows(data, care_id)
        for sizes, page_label in (({"슈싱", "퀸"}, "슈싱, 퀸"), ({"킹", "라지킹"}, "킹, 라지킹")):
            add_detail_page(doc, care_id, rows, sizes, page_label, start_new_page=not first_page)
            first_page = False
    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / "비렉스_5년_케어옵션_전체조합상세표_8장_인쇄용_7월15프로.docx"
    doc.save(path)
    audit_detail_collection(path)
    return path


def main():
    data = load_product_data()
    for care_id in ("serviceFree", "basic", "special", "total"):
        path, rows, summary = build_doc(care_id, data)
        print(path)
        print(
            f"{care_id}: rows={len(rows)} higher={summary['higher_count']} "
            f"lower={summary['lower_count']} avg_diff={summary['avg_diff']:.0f} "
            f"min_diff={min(r['diff'] for r in rows):.0f} max_diff={max(r['diff'] for r in rows):.0f}"
        )
    print(build_detail_collection_doc(data))


if __name__ == "__main__":
    main()
