import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path("index.html")
OUT_DIR = Path("proposal_output")
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "제휴카드할인_기본할인_추가할인_고객제안서_A4.docx"

FONT = "Malgun Gothic"
INK = "172033"
BLUE = "1F5AA6"
DARK_BLUE = "173B66"
RED = "C4312E"
GRAY = "5F6B7A"
LIGHT_BLUE = "EAF2FF"
LIGHT_RED = "FFF1F0"
PALE_GOLD = "FFF4D6"
BORDER = "C8D1E0"


def fmt_won(value):
    return f"{value:,}원"


def read_card_data():
    source = SOURCE.read_text(encoding="utf-8")
    block_match = re.search(r"const CARD_PROMO_DATA = \[(.*?)\n\s*\];", source, re.S)
    if not block_match:
        raise RuntimeError("CARD_PROMO_DATA block not found")
    block = block_match.group(1)
    cards = []
    pattern = re.compile(
        r"\{\s*id:'(?P<id>[^']+)',\s*name:'(?P<name>[^']+)'[\s\S]*?"
        r"tiers:\[(?P<tiers>[\s\S]*?)\]\s*,\s*period:(?P<period>\d+|0),\s*tel:'(?P<tel>[^']*)'[\s\S]*?"
        r"note:'(?P<note>[^']*)'\s*\}",
        re.S,
    )
    tier_pattern = re.compile(r"\{spend:(\d+),\s*base:(\d+),\s*promo:(\d+),\s*total:(\d+)\}")
    for match in pattern.finditer(block):
        tiers = []
        for tier in tier_pattern.finditer(match.group("tiers")):
            spend, base, promo, total = map(int, tier.groups())
            if base + promo != total:
                raise RuntimeError(f"{match.group('name')} {spend} tier total mismatch")
            tiers.append({"spend": spend, "base": base, "promo": promo, "total": total})
        if tiers:
            cards.append(
                {
                    "id": match.group("id"),
                    "name": match.group("name"),
                    "period": int(match.group("period")),
                    "tel": match.group("tel"),
                    "note": match.group("note"),
                    "tiers": tiers,
                }
            )
    if not cards:
        raise RuntimeError("No card tiers parsed")
    return cards


def set_run_font(run, size=None, bold=False, color=INK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_para_style(p, before=0, after=3, line=1.0, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align:
        p.alignment = align


def add_text(p, text, size=9.3, bold=False, color=INK):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return r


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear_cell(cell):
    for p in cell.paragraphs:
        p.clear()


def fill_cell(cell, text, size=8.4, bold=False, color=INK, align=None, fill=None):
    clear_cell(cell)
    p = cell.paragraphs[0]
    set_para_style(p, after=0, line=1.0, align=align)
    add_text(p, text, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)
    if fill:
        shade_cell(cell, fill)


def set_table_width(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")


cards = read_card_data()
all_tiers = [(card, tier) for card in cards for tier in card["tiers"]]
max_total_card, max_total_tier = max(all_tiers, key=lambda item: item[1]["total"])
max_promo_card, max_promo_tier = max(all_tiers, key=lambda item: item[1]["promo"])
promo_cards = [card for card in cards if card["period"] > 0]

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.orientation = WD_ORIENT.PORTRAIT
section.top_margin = Cm(0.9)
section.bottom_margin = Cm(0.85)
section.left_margin = Cm(1.05)
section.right_margin = Cm(1.05)
section.header_distance = Cm(0.5)
section.footer_distance = Cm(0.5)

styles = doc.styles
styles["Normal"].font.name = FONT
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Normal"].font.size = Pt(9.2)

title = doc.add_paragraph()
set_para_style(title, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(title, "제휴카드 할인 고객 제안서", size=19, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
set_para_style(subtitle, after=7, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(subtitle, "앱 반영 기준: 기본할인 + 추가할인 금액을 오늘 현장에서 바로 비교해 드립니다.", size=9.4, bold=True, color=GRAY)

hero = doc.add_table(rows=1, cols=3)
set_table_width(hero, [6.2, 6.2, 6.5])
hero_items = [
    ("최대 월 혜택", f"{max_total_card['name']} {max_total_tier['spend']}만↑", fmt_won(max_total_tier["total"])),
    ("추가할인 최대", f"{max_promo_card['name']} {max_promo_tier['spend']}만↑", fmt_won(max_promo_tier["promo"])),
    ("추가할인 기간", f"{len(promo_cards)}개 카드 적용", "최대 60개월"),
]
for i, (label, detail, value) in enumerate(hero_items):
    cell = hero.cell(0, i)
    fill = LIGHT_RED if i == 1 else LIGHT_BLUE
    shade_cell(cell, fill)
    set_cell_border(cell, color="AFC5E8" if i != 1 else "F0B8B5", size="10")
    set_cell_margins(cell, top=120, bottom=120, start=120, end=120)
    clear_cell(cell)
    p = cell.paragraphs[0]
    set_para_style(p, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, label, size=8.2, bold=True, color=GRAY)
    p = cell.add_paragraph()
    set_para_style(p, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, value, size=14.2, bold=True, color=RED if i == 1 else DARK_BLUE)
    p = cell.add_paragraph()
    set_para_style(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, detail, size=8.0, bold=True, color=INK)

p = doc.add_paragraph()
set_para_style(p, before=6, after=4, line=1.08)
add_text(
    p,
    "제휴카드는 상시 적용되는 기본할인과 프로모션 기간에 붙는 추가할인이 합산되어 월 혜택이 결정됩니다. 고객님 카드 사용 실적에 맞는 구간을 오늘 확인하면, 기본 할인만 안내받는 것보다 더 큰 월 절감 조건을 바로 비교할 수 있습니다.",
    size=9.2,
)

section_title = doc.add_paragraph()
set_para_style(section_title, before=2, after=3)
add_text(section_title, "현재 앱 반영 제휴카드 할인표", size=11.6, bold=True, color=DARK_BLUE)

table = doc.add_table(rows=1, cols=5)
set_table_width(table, [3.0, 3.3, 3.2, 3.2, 6.2])
headers = ["카드", "최대혜택 구간", "기본할인", "추가할인", "월 총혜택 / 기간"]
for idx, header in enumerate(headers):
    fill_cell(table.rows[0].cells[idx], header, size=8.0, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, fill=DARK_BLUE)

for card in cards:
    max_tier = max(card["tiers"], key=lambda tier: tier["total"])
    range_text = f"{max_tier['spend']}만↑"
    base_text = fmt_won(max_tier["base"])
    promo_text = "없음" if max_tier["promo"] == 0 else fmt_won(max_tier["promo"])
    total_text = fmt_won(max_tier["total"])
    if card["period"] > 0:
        point = f"{total_text} / 추가 {card['period']}개월"
    else:
        point = f"{total_text} / 기본할인형"
    cells = table.add_row().cells
    values = [card["name"], range_text, base_text, promo_text, point]
    for idx, value in enumerate(values):
        color = RED if idx == 3 and value != "없음" else INK
        bold = idx in (0, 3, 4)
        fill = "FFF8F7" if idx == 3 and value != "없음" else None
        fill_cell(cells[idx], value, size=7.65, bold=bold, color=color, align=WD_ALIGN_PARAGRAPH.CENTER if idx < 4 else None, fill=fill)

closing = doc.add_table(rows=1, cols=1)
cell = closing.cell(0, 0)
shade_cell(cell, PALE_GOLD)
set_cell_border(cell, color="E6C45A", size="10")
set_cell_margins(cell, top=110, bottom=110, start=150, end=150)
clear_cell(cell)
p = cell.paragraphs[0]
set_para_style(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "오늘 상담에서 바로 확인해야 하는 이유", size=11.2, bold=True, color="7A5200")
p = cell.add_paragraph()
set_para_style(p, after=0, line=1.07, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(
    p,
    "기본할인은 카드별로 고정되어 있지만, 추가할인은 카드·실적·신청 시점·프로모션 기간에 따라 달라집니다. 지금 상담 중 고객님 조건에 맞는 카드 구간을 잡으면 월 납입액을 바로 낮춰서 제안드릴 수 있습니다.",
    size=9.1,
    bold=True,
)

fine = doc.add_paragraph()
set_para_style(fine, before=4, after=1, line=1.0)
add_text(fine, "안내 사항  ", size=7.6, bold=True, color=GRAY)
add_text(
    fine,
    "위 금액은 현재 앱의 제휴카드 데이터 기준입니다. 실제 적용은 카드사 심사, 전월 실적, 자동이체 신청, 제품 계약 조건, 프로모션 운영 시점에 따라 달라질 수 있으며 최종 혜택은 상담 후 확정됩니다.",
    size=7.4,
    color=GRAY,
)

footer = doc.add_paragraph()
set_para_style(footer, before=2, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
add_text(footer, "상담 담당자: ____________________   연락처: ____________________", size=8.2, bold=True, color=DARK_BLUE)

doc.save(OUT_PATH)
print(OUT_PATH)
